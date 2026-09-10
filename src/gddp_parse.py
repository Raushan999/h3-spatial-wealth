"""Parsers that turn official DES publications into district GDDP records.

Every parser reads a file that is already persisted under `data/raw/gddp/{STATE}/`
and returns `(records, provenance)` where a record is
`{"district": <name as published>, "y_crore": float, "source_name": str}`.

Nothing here fabricates a value: if a table cannot be located the parser raises.
"""

from __future__ import annotations

import re
from pathlib import Path

from .config import GddpSource


def _clean(text) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _to_float(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"[^\d.\-]", "", str(value))
    if text in {"", "-", "."}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _load_sheet(path: Path, sheet: str | None):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    name = sheet if sheet in wb.sheetnames else wb.sheetnames[0]
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    wb.close()
    return name, rows


def parse_xlsx_matrix(path: Path, src: GddpSource) -> tuple[list[dict], dict]:
    """DES sheet with districts across columns and economic activities down rows (UP)."""
    sheet, rows = _load_sheet(path, src.sheet)
    if not rows:
        raise RuntimeError(f"{path} sheet {sheet!r} is empty")

    # The district header is the early row with the most short text labels.
    header_idx, best = None, 0
    for i, row in enumerate(rows[:6]):
        n = sum(
            1
            for c in row[2:]
            if isinstance(c, str) and 0 < len(_clean(c)) <= 40 and not _clean(c).isdigit()
        )
        if n > best:
            header_idx, best = i, n
    if header_idx is None or best < 5:
        raise RuntimeError(f"No district header row found in {path}")

    label = _clean(src.row_label or "").lower()
    value_idx = None
    for i, row in enumerate(rows):
        if i <= header_idx:
            continue
        if label and label in _clean(row[1] if len(row) > 1 else "").lower():
            value_idx = i  # keep the last match: totals sit below sub-totals
    if value_idx is None:
        raise RuntimeError(f"Row {src.row_label!r} not found in {path} sheet {sheet!r}")

    drop = {d.lower() for d in src.drop_names}
    header = rows[header_idx]
    values = rows[value_idx]
    records: list[dict] = []
    for col in range(2, len(header)):
        name = _clean(header[col])
        if not name or name.lower() in drop or "region" in name.lower():
            continue
        y = _to_float(values[col] if col < len(values) else None)
        if y is None or y <= 0:
            continue
        records.append(
            {
                "district": name,
                "y_crore": y * src.value_scale,
                "source_name": f"{name} — {_clean(values[1])}",
            }
        )
    if not records:
        raise RuntimeError(f"No district values parsed from {path}")
    return records, {
        "parser": "xlsx_matrix",
        "sheet": sheet,
        "header_row": header_idx + 1,
        "value_row": value_idx + 1,
        "value_row_label": _clean(values[1]),
        "n_districts": len(records),
    }


def parse_xlsx_table(path: Path, src: GddpSource) -> tuple[list[dict], dict]:
    """Simple sheet with one district per row (Bihar)."""
    sheet, rows = _load_sheet(path, src.sheet)
    want_name = _clean(src.district_field).lower()
    want_val = _clean(src.value_field).lower()
    header_idx = name_col = val_col = None
    for i, row in enumerate(rows[:12]):
        cells = [_clean(c).lower() for c in row]
        if want_name in cells and want_val in cells:
            header_idx, name_col, val_col = i, cells.index(want_name), cells.index(want_val)
            break
    if header_idx is None:
        raise RuntimeError(
            f"Header with {src.district_field!r} and {src.value_field!r} not found in {path}"
        )

    drop = {d.lower() for d in src.drop_names}
    records: list[dict] = []
    for row in rows[header_idx + 1 :]:
        name = _clean(row[name_col] if name_col < len(row) else "")
        if not name or name.lower() in drop:
            continue
        y = _to_float(row[val_col] if val_col < len(row) else None)
        if y is None or y <= 0:
            continue
        records.append(
            {"district": name, "y_crore": y * src.value_scale, "source_name": src.value_field}
        )
    if not records:
        raise RuntimeError(f"No district rows parsed from {path}")
    title = _clean(rows[0][0] if rows and rows[0] else "")
    return records, {
        "parser": "xlsx_table",
        "sheet": sheet,
        "title": title,
        "value_column": src.value_field,
        "n_districts": len(records),
    }


_ROW_RE = re.compile(
    r"^\s*(?P<sn>\d{1,3})?\s*(?P<name>[A-Za-z][A-Za-z.\s'()\-]*?)\s+(?P<nums>[\d,]+(?:\.\d+)?(?:\s+[\d,]+(?:\.\d+)?)*)\s*$"
)


def parse_pdf_table(path: Path, src: GddpSource) -> tuple[list[dict], dict]:
    """District x year table in a text-extractable DES PDF (Rajasthan Table 3)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    label = _clean(src.row_label or "").lower()
    if not label:
        raise RuntimeError("pdf_table needs row_label set to the table caption")

    drop = {d.lower() for d in src.drop_names}
    records: dict[str, dict] = {}
    used_pages: list[int] = []
    for page_no, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if label not in _clean(text).lower():
            continue
        lines = [ln for ln in text.splitlines() if ln.strip()]
        # Year header: the line holding the fiscal-year column names.
        years: list[str] = []
        for ln in lines:
            found = re.findall(r"\d{4}-\d{2}", ln)
            if len(found) >= 2:
                years = found
                break
        if src.year not in years:
            continue
        col = years.index(src.year)
        for ln in lines:
            m = _ROW_RE.match(ln)
            if not m:
                continue
            name = _clean(m.group("name"))
            if not name or name.lower() in drop or len(name) < 3:
                continue
            nums = [_to_float(x) for x in m.group("nums").split()]
            if len(nums) != len(years) or nums[col] is None or nums[col] <= 0:
                continue
            records[name.lower()] = {
                "district": name,
                "y_crore": nums[col] * src.value_scale,
                "source_name": f"{_clean(src.row_label)} [{src.year}]",
            }
            if page_no not in used_pages:
                used_pages.append(page_no)
    if not records:
        raise RuntimeError(f"Table {src.row_label!r} for {src.year} not parsed from {path}")
    return list(records.values()), {
        "parser": "pdf_table",
        "table": _clean(src.row_label),
        "year_column": src.year,
        "pages": [p + 1 for p in used_pages],
        "published_units": src.units if src.value_scale == 1.0 else "INR lakh (rescaled to crore)",
        "n_districts": len(records),
    }


def parse_gsdp_single_pdf(path: Path, src: GddpSource, unit_name: str) -> tuple[list[dict], dict]:
    """State/UT-wide GSDP quoted in an official abstract (Delhi NCT)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    label = _clean(src.row_label or "GSDP at constant").lower()
    basis = "constant" if src.price_basis == "constant" else "current"
    for page_no, page in enumerate(reader.pages):
        flat = _clean(page.extract_text() or "")
        low = flat.lower()
        if label not in low or f"{basis}" not in low:
            continue
        window = flat[low.index(label) :][:600]
        hits = re.findall(r"([\d,]{4,})\s*Crore in (\d{4}-\d{2})", window)
        value = next((v for v, yr in hits if yr == src.year), None)
        if value is None:
            continue
        y = _to_float(value)
        if not y:
            continue
        return (
            [
                {
                    "district": unit_name,
                    "y_crore": y * src.value_scale,
                    "source_name": (
                        f"GSDP at {basis} ({src.base_year}) prices {src.year}, "
                        "Statistical Abstract of Delhi 2024 executive summary"
                    ),
                }
            ],
            {
                "parser": "gsdp_single_pdf",
                "page": page_no + 1,
                "quote": window[:220],
                "geography": "Delhi / NCT only (not NCR)",
                "n_districts": 1,
            },
        )
    raise RuntimeError(
        f"Could not read {basis}-price GSDP for {src.year} from {path}. "
        "Refusing to fall back to a hardcoded figure."
    )


def parse_official_source(state_code: str, path: Path, src: GddpSource, unit_name: str):
    if src.kind == "xlsx_matrix":
        return parse_xlsx_matrix(path, src)
    if src.kind == "xlsx_table":
        return parse_xlsx_table(path, src)
    if src.kind == "pdf_table":
        return parse_pdf_table(path, src)
    if src.kind == "gsdp_single":
        return parse_gsdp_single_pdf(path, src, unit_name)
    raise RuntimeError(f"No parser for GDDP kind {src.kind!r} ({state_code})")
